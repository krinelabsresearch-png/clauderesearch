"""Ricalcola modello-cassa.xlsx con un valutatore di formule indipendente e lo confronta con motore.py."""
import sys, time, formulas
from motore import scenario, simula, valore_cliente, SCENARI
t = time.time()
xl = formulas.ExcelModel().loads("modello-cassa.xlsx").finish()
sol = xl.calculate()
print("calcolato in", round(time.time() - t), "s", file=sys.stderr)
def val(sheet, cell):
    for k, v in sol.items():
        if k.upper().endswith(f"{sheet.upper()}'!{cell.upper()}"):
            x = v.value[0, 0] if hasattr(v, "value") else v
            return float(x)
    raise KeyError(sheet + cell)
import openpyxl
from openpyxl.utils import get_column_letter as L
wb = openpyxl.load_workbook("modello-cassa.xlsx")
SHEETS = ["S1 Base", "S2 Ritardo", "S3 Avverso", "S4 Avverso+arresto", "S5 Tutto insieme"]
maxdiff = 0; report = []
for nome, sh in zip(SCENARI, SHEETS):
    ws = wb[sh]
    rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    r = simula(scenario(nome))
    checks = {"Cassa a fine mese": r["cassa"], "Incassi lordi (IVA incl.)": r["R"]["incassi"],
              "Rimborsi di garanzia pagati (lordi)": r["R"]["rimborsi_lordi"], "IVA versata nel mese": r["iva_da_versare"],
              "Credito IVA a fine mese": r["credito"], "Flaconi in magazzino a fine mese": r["scorta"],
              "Nuovi clienti totali": r["nuovi"], "Garanzia: rimborsi attesi non ancora pagati (lordi)": r["esposizione"],
              "Risultato operativo del mese (competenza semplificata)": r["operativo"]}
    d_sc = 0
    for lab, serie in checks.items():
        rr = rows[lab]
        for m in range(1, 31):
            x = val(sh, f"{L(2+m)}{rr}")
            d = abs(x - serie[m]); d_sc = max(d_sc, d)
            if d > 0.01: report.append((nome, lab, m, x, serie[m]))
    s = scenario(nome)
    for h, cell in ((12, "Valore del cliente medio a 12 mesi"), (24, "Valore del cliente medio a 24 mesi")):
        x = val(sh, f"C{rows[cell]}"); y = valore_cliente(s, h)["medio"]
        d_sc = max(d_sc, abs(x - y))
        if abs(x - y) > 0.01: report.append((nome, cell, h, x, y))
    print(f"{nome:24s} differenza massima {d_sc:.6f}  cassa minima foglio {val(sh, 'C'+str(rows['Cassa minima'])):,.0f}")
    maxdiff = max(maxdiff, d_sc)
print("DIFFERENZA MASSIMA COMPLESSIVA", maxdiff)
for x in report[:30]: print(x)
